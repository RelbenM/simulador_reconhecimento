"""Simulador de Reconhecimento — Gestão de Pessoas (GP).

Execução: python app.py
"""
import io
import math
import os
import re
import threading
from datetime import datetime
import unicodedata
import webbrowser

from flask import Flask, abort, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import database as db

app = Flask(__name__)
OUTPUT_DIR = os.path.join(db.BASE_DIR, "output")

CAMPOS_COLAB = ["maturidade", "conduta", "zona", "elegivel"]
TAM_MAX_TEXTO = 100
ROTULOS = {"maturidade": "Maturidade", "conduta": "Conduta", "zona": "Zona salarial",
           "elegivel": "Elegível", "salario": "Salário atual"}


# ----------------------------------------------------------------- formatação
def brl(v):
    if v is None:
        return "—"
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def num_br(v):
    if v is None:
        return ""
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct_num(v):
    """0.5 -> '50'; 0.455 -> '45,5' (para campos de entrada)."""
    if v is None:
        return ""
    n = round(v * 100, 2)
    return str(int(n)) if n == int(n) else f"{n:.2f}".rstrip("0").replace(".", ",")


def pct(v):
    return f"{pct_num(v)}%" if v is not None else "—"


def tempo(meses):
    anos, resto = divmod(int(meses), 12)
    m = f"{resto} {'mês' if resto == 1 else 'meses'}"
    if anos == 0:
        return m
    a = f"{anos} {'ano' if anos == 1 else 'anos'}"
    return a if resto == 0 else f"{a} e {m}"


def datahora(iso):
    return datetime.fromisoformat(iso).strftime("%d/%m/%Y %H:%M:%S")


app.jinja_env.filters.update(brl=brl, num_br=num_br, pct=pct, pct_num=pct_num, tempo=tempo,
                             datahora=datahora)


@app.context_processor
def contexto_global():
    ativo = request.view_args.get("setor_id") if request.endpoint == "setor" else None
    return {"setores_menu": db.listar_setores_menu(), "setor_ativo_id": ativo,
            "STATUS": db.STATUS, "CICLO": db.CICLO}


# ------------------------------------------------------------------ validação
def _num(valor):
    """Converte texto/número em float. None se vazio; ValueError se inválido."""
    if valor is None or isinstance(valor, bool):
        if valor is None:
            return None
        raise ValueError
    if isinstance(valor, (int, float)):
        n = float(valor)
    else:
        s = str(valor).strip().replace("R$", "").replace("%", "").replace(" ", "")
        if s == "":
            return None
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
        n = float(s)
    if not math.isfinite(n):
        raise ValueError
    return n


def _num_ou_erro(valor):
    try:
        return _num(valor)
    except ValueError:
        raise ValueError("Valor inválido.")


def parse_orcamento(valor):
    n = _num_ou_erro(valor)
    if n is None:
        raise ValueError("Informe o orçamento da folha.")
    if n < 0:
        raise ValueError("O orçamento da folha não pode ser negativo.")
    return round(n, 2)


def parse_percentual(valor):
    n = _num_ou_erro(valor)
    if n is None:
        raise ValueError("Informe o percentual de reconhecimento.")
    if n < 0 or n > 100:
        raise ValueError("O percentual deve estar entre 0% e 100%.")
    return round(n / 100, 4)


def validar(setor_id, dados):
    """Retorna (limpo, erros, pendencias). Erros impedem salvar; pendências não."""
    erros, pendencias = [], []
    limpo = {"orcamento": None, "percentual": None, "colaboradores": []}

    try:
        limpo["orcamento"] = parse_orcamento(dados.get("orcamento"))
    except ValueError as e:
        msg = str(e) if str(e) else "Orçamento inválido."
        erros.append({"campo": "orcamento", "mensagem": msg})
    try:
        limpo["percentual"] = parse_percentual(dados.get("percentual"))
    except ValueError as e:
        msg = str(e) if str(e) else "Percentual inválido."
        erros.append({"campo": "percentual", "mensagem": msg})

    enviados = {}
    for item in dados.get("colaboradores") or []:
        try:
            enviados[int(item.get("id"))] = item
        except (TypeError, ValueError):
            continue

    for c in db.listar_colaboradores(setor_id):
        item = enviados.get(c["id"], {})
        novo = {"id": c["id"]}
        faltando = []
        for campo in CAMPOS_COLAB:
            valor = (item.get(campo, c[campo]) or "").strip() or None
            if valor is not None and campo in db.OPCOES and valor not in db.OPCOES[campo]:
                erros.append({"campo": f"colab-{c['id']}-{campo}",
                              "mensagem": f"{c['nome']}: valor inválido em {ROTULOS[campo]}."})
                valor = None
            elif valor is not None and len(valor) > TAM_MAX_TEXTO:
                erros.append({"campo": f"colab-{c['id']}-{campo}",
                              "mensagem": f"{c['nome']}: {ROTULOS[campo]} excede {TAM_MAX_TEXTO} caracteres."})
                valor = None
            elif valor is None:
                faltando.append(campo)
            novo[campo] = valor
        try:
            sal = _num(item["salario"]) if "salario" in item else c["salario"]
            if sal is not None and sal < 0:
                raise ValueError
            novo["salario"] = round(sal, 2) if sal is not None else None
            if sal is None:
                faltando.append("salario")
        except ValueError:
            erros.append({"campo": f"colab-{c['id']}-salario",
                          "mensagem": f"{c['nome']}: salário inválido ou negativo."})
            novo["salario"] = None
        for campo in faltando:
            pendencias.append({"campo": f"colab-{c['id']}-{campo}",
                               "mensagem": f"{c['nome']}: preencher {ROTULOS[campo]}."})
        limpo["colaboradores"].append(novo)
    return limpo, erros, pendencias


def definir_status(limpo, pendencias, enviar=False):
    if enviar:
        return "enviado"
    nada = all(not c[f] for c in limpo["colaboradores"] for f in CAMPOS_COLAB)
    if nada:
        return "nao_iniciado"
    return "em_preenchimento" if pendencias else "pronto_envio"


def resposta_dados(setor_id, limpo, status, mensagem, pendencias=None):
    return {
        "ok": True, "mensagem": mensagem, "status": status, "status_label": db.STATUS[status],
        "orcamento": brl(limpo["orcamento"]), "percentual": pct(limpo["percentual"]),
        "disponivel": brl(round(limpo["orcamento"] * limpo["percentual"], 2)),
        "pendencias": pendencias or [],
    }


# ---------------------------------------------------------------------- rotas
@app.route("/")
def home():
    return render_template("home.html", setores=db.listar_setores(), totais=db.totais())


@app.route("/setores/<int:setor_id>")
def setor(setor_id):
    s = db.obter_setor(setor_id)
    if not s:
        abort(404)
    return render_template("setor.html", setor=s, colaboradores=db.listar_colaboradores(setor_id),
                           opcoes=db.OPCOES)


@app.post("/api/calcular")
def api_calcular():
    """Cálculo definitivo do valor disponível (a tela apenas exibe o resultado)."""
    dados = request.get_json(silent=True) or {}
    erros = {}
    orc = pc = None
    try:
        orc = parse_orcamento(dados.get("orcamento"))
    except ValueError as e:
        erros["orcamento"] = str(e) or "Orçamento inválido."
    try:
        pc = parse_percentual(dados.get("percentual"))
    except ValueError as e:
        erros["percentual"] = str(e) or "Percentual inválido."
    if erros:
        return jsonify(ok=False, erros=erros), 422
    return jsonify(ok=True, disponivel=brl(round(orc * pc, 2)))


def _processar(setor_id, enviar):
    if not db.obter_setor(setor_id):
        abort(404)
    dados = request.get_json(silent=True) or {}
    limpo, erros, pendencias = validar(setor_id, dados)
    if erros:
        return jsonify(ok=False, mensagem="Corrija os campos destacados. Nada foi perdido.",
                       erros=erros), 422
    if enviar and pendencias:
        return jsonify(ok=False, mensagem="Existem campos obrigatórios pendentes. "
                       "Preencha-os para enviar.", erros=pendencias), 422
    status = definir_status(limpo, pendencias, enviar)
    db.salvar_setor(setor_id, limpo["orcamento"], limpo["percentual"], limpo["colaboradores"], status)
    if enviar:
        db.registrar_log("Envio ao gerente", db.obter_setor(setor_id)["gerente"],
                         limpo["orcamento"], limpo["percentual"])
        msg = "Dados validados e salvos. Setor marcado como enviado (sem envio real de e-mail)."
    elif pendencias:
        msg = "Alterações salvas. Ainda há campos pendentes."
    else:
        msg = "Alterações salvas com sucesso. Setor pronto para envio."
    return jsonify(resposta_dados(setor_id, limpo, status, msg, pendencias))


@app.post("/api/setores/<int:setor_id>/salvar")
def api_salvar(setor_id):
    return _processar(setor_id, enviar=False)


@app.post("/api/setores/<int:setor_id>/enviar")
def api_enviar(setor_id):
    return _processar(setor_id, enviar=True)


def _slug(texto):
    s = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")


@app.route("/setores/<int:setor_id>/excel")
def excel(setor_id):
    s = db.obter_setor(setor_id)
    if not s:
        abort(404)
    _, erros, pendencias = validar(setor_id, {"orcamento": s["orcamento"], "percentual": s["percentual"]})
    if erros or pendencias:
        return jsonify(ok=False, mensagem="Não é possível gerar o Excel: há campos pendentes de resposta.",
                       erros=erros + pendencias), 422
    colabs = db.listar_colaboradores(setor_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Simulação"
    verde, verde_claro = "087A4B", "E8F5EE"
    fino = Side(style="thin", color="D5DEE3")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    moeda = '"R$" #,##0.00'

    ws["A1"] = "Simulador de Reconhecimento — " + s["nome"]
    ws["A1"].font = Font(bold=True, size=14, color="004D36")
    resumo = [("Setor", s["nome"], None), ("Orçamento da folha", s["orcamento"], moeda),
              ("Percentual de reconhecimento", s["percentual"], "0.0%"),
              ("Valor disponível", s["disponivel"], moeda)]
    for i, (rot, val, fmt) in enumerate(resumo, start=3):
        ws.cell(i, 1, rot).font = Font(bold=True)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=verde_claro)
        c = ws.cell(i, 2, val)
        c.alignment = Alignment(horizontal="left")
        if fmt:
            c.number_format = fmt

    cab = ["Colaborador", "Maturidade", "Tempo de cargo", "Conduta", "Salário atual",
           "Zona salarial", "Elegível"]
    linha_cab = 8
    for j, nome in enumerate(cab, start=1):
        c = ws.cell(linha_cab, j, nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=verde)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borda
    for i, c in enumerate(colabs, start=linha_cab + 1):
        valores = [c["nome"], c["maturidade"], tempo(c["tempo_cargo"]), c["conduta"],
                   c["salario"], c["zona"], c["elegivel"]]
        for j, v in enumerate(valores, start=1):
            cel = ws.cell(i, j, v)
            cel.border = borda
            if j == 5:
                cel.number_format = moeda
    ultima = linha_cab + max(len(colabs), 1)
    ws.auto_filter.ref = f"A{linha_cab}:{get_column_letter(len(cab))}{ultima}"
    ws.freeze_panes = ws.cell(linha_cab + 1, 1)
    for j, larg in enumerate([32, 18, 20, 22, 18, 16, 12], start=1):
        ws.column_dimensions[get_column_letter(j)].width = larg
    ws.column_dimensions["B"].width = 30  # comporta valores do resumo

    nome_arquivo = f"Simulacao_Reconhecimento_{_slug(s['nome'])}.xlsx"
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb.save(os.path.join(OUTPUT_DIR, nome_arquivo))
    db.registrar_log("Excel gerado", s["gerente"], s["orcamento"], s["percentual"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=nome_arquivo,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/home/parametros")
def api_home_parametros():
    """Homepage: pré-visualiza (salvar=false) ou grava (salvar=true) folha e percentual dos setores."""
    dados = request.get_json(silent=True) or {}
    salvar = bool(dados.get("salvar"))
    enviados = {}
    for it in dados.get("setores") or []:
        try:
            enviados[int(it.get("id"))] = it
        except (TypeError, ValueError):
            continue
    erros, linhas = [], []
    for s in db.listar_setores():
        it = enviados.get(s["id"], {})
        orc, pc = s["orcamento"], s["percentual"]
        if "orcamento" in it:
            try:
                orc = parse_orcamento(it["orcamento"])
            except ValueError as e:
                erros.append({"campo": f"orcamento-{s['id']}", "mensagem": f"{s['nome']}: {e}"})
                orc = s["orcamento"]
        if "percentual" in it:
            try:
                pc = parse_percentual(it["percentual"])
            except ValueError as e:
                erros.append({"campo": f"percentual-{s['id']}", "mensagem": f"{s['nome']}: {e}"})
                pc = s["percentual"]
        linhas.append({"id": s["id"], "orcamento": orc, "percentual": pc,
                       "disponivel": round(orc * pc, 2), "qtd": s["qtd_colaboradores"]})
    if salvar and erros:
        return jsonify(ok=False, mensagem="Corrija os campos destacados. Nada foi perdido.", erros=erros), 422
    if salvar:
        db.salvar_parametros(linhas)
    return jsonify(
        ok=True, erros=erros,
        mensagem="Orçamento e percentuais salvos com sucesso." if salvar else "",
        linhas={str(l["id"]): brl(l["disponivel"]) for l in linhas},
        totais={"colaboradores": sum(l["qtd"] for l in linhas),
                "folha": brl(round(sum(l["orcamento"] for l in linhas), 2)),
                "disponivel": brl(round(sum(l["disponivel"] for l in linhas), 2))},
    )


@app.route("/logs")
def logs():
    return render_template("logs.html", registros=db.listar_logs())


def abrir_navegador(url):
    threading.Timer(1.0, lambda: webbrowser.open(url)).start()


if __name__ == "__main__":
    db.init_db()
    host, porta = "127.0.0.1", 5000
    url = f"http://{host}:{porta}"
    print(f"\n  Simulador de Reconhecimento rodando em: {url}\n  (Ctrl+C para encerrar)\n")
    abrir_navegador(url)
    app.run(host=host, port=porta, debug=False)
